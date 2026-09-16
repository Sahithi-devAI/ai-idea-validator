"""
Runs the validator against all ideas in data/product_hunt_ideas.xlsx,
writes dimension scores, build signal, verdict, and mismatch back into the file.
Saves after each row — safe to re-run if interrupted.
"""

import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

import openpyxl
from dotenv import load_dotenv

from core.logger import log_evaluation
from main import run_evaluation

load_dotenv()

EXCEL_PATH = os.path.join(os.path.dirname(__file__), "..", "data", "product_hunt_ideas.xlsx")

NEW_VALIDATOR_COLS = [
    "validator_score_market",
    "validator_score_usefulness",
    "validator_score_feasibility",
    "validator_score_monetization",
    "validator_build_signal",
    "validator_verdict",
]

DIMENSION_MAP = {
    "validator_score_market":       "Market Potential",
    "validator_score_usefulness":   "Product Usefulness",
    "validator_score_feasibility":  "Feasibility & Defensibility",
    "validator_score_monetization": "Monetization Viability",
}


def restructure_columns(ws):
    headers = [cell.value for cell in ws[1]]
    if "validator_score_market" in headers:
        print("Columns already restructured — skipping.")
        return

    old_col = headers.index("validator_output") + 1  # 1-indexed
    ws.delete_cols(old_col)
    ws.insert_cols(old_col, len(NEW_VALIDATOR_COLS))

    for i, name in enumerate(NEW_VALIDATOR_COLS):
        ws.cell(row=1, column=old_col + i).value = name

    print("Columns restructured.")


def col_index(ws, name):
    headers = [cell.value for cell in ws[1]]
    return headers.index(name) + 1  # 1-indexed


def main():
    provider = os.getenv("LLM_PROVIDER", "anthropic")
    prompt_version = os.getenv("PROMPT_VERSION", "v1_5")

    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb.active

    restructure_columns(ws)
    wb.save(EXCEL_PATH)

    total = ws.max_row - 1
    passed = skipped = errors = 0

    for row_idx in range(2, ws.max_row + 1):
        idea_id      = ws.cell(row=row_idx, column=col_index(ws, "idea_id")).value
        idea_summary = ws.cell(row=row_idx, column=col_index(ws, "idea_summary")).value
        human_signal = ws.cell(row=row_idx, column=col_index(ws, "human_expected_signal")).value

        # Skip rows already evaluated
        if ws.cell(row=row_idx, column=col_index(ws, "validator_build_signal")).value:
            print(f"[{idea_id}] Already evaluated — skipping.")
            skipped += 1
            continue

        print(f"[{idea_id}] Evaluating... ", end="", flush=True)

        try:
            result = run_evaluation(idea_summary, provider=provider, prompt_version=prompt_version)
            log_evaluation(result, provider=provider, prompt_version=prompt_version)

            scores = {d.name: d.score for d in result.dimensions}

            for col_name, dim_name in DIMENSION_MAP.items():
                ws.cell(row=row_idx, column=col_index(ws, col_name)).value = scores.get(dim_name)

            ws.cell(row=row_idx, column=col_index(ws, "validator_build_signal")).value = result.build_signal.value
            ws.cell(row=row_idx, column=col_index(ws, "validator_verdict")).value = result.verdict

            mismatch = result.build_signal.value != human_signal if human_signal else None
            ws.cell(row=row_idx, column=col_index(ws, "validator_mismatch")).value = mismatch

            wb.save(EXCEL_PATH)

            status = "MISMATCH" if mismatch else "match"
            print(f"{result.build_signal.value} (expected: {human_signal}) — {status}")
            passed += 1

        except Exception as e:
            print(f"ERROR — {e}")
            errors += 1

    print(f"\n{'=' * 50}")
    print(f"  Done: {passed} evaluated | {skipped} skipped | {errors} errors")
    print(f"  Results saved to {EXCEL_PATH}")
    print(f"{'=' * 50}")


if __name__ == "__main__":
    main()
